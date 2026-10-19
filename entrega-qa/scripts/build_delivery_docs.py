"""Regenerate the three main hand-in Word documents from Markdown sources."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentos"
JOBS = [
    ("02-test-plan.md", "Plan_de_Pruebas_Hospital.docx", False),
    ("03-historias-y-test-cases.md", "Casos_de_Prueba_Hospital.docx", False),
    ("01-scrum-y-reparto.md", "Organizacion_Equipo_Hospital.docx", True),
]


def set_font(style, size: float, bold: bool = False) -> None:
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")


def configure(doc: Document, landscape: bool) -> None:
    sec = doc.sections[0]
    sec.page_width = Inches(11 if landscape else 8.5)
    sec.page_height = Inches(8.5 if landscape else 11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    styles = doc.styles
    set_font(styles["Normal"], 10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    set_font(styles["Title"], 18, True)
    styles["Title"].paragraph_format.space_after = Pt(12)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is None:
        title_border = OxmlElement("w:pBdr")
        title_ppr.append(title_border)
    bottom = title_border.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        title_border.append(bottom)
    bottom.set(qn("w:val"), "nil")
    set_font(styles["Heading 1"], 13, True)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    set_font(styles["Heading 2"], 11.5, True)
    styles["Heading 2"].paragraph_format.space_before = Pt(11)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    for list_style in ("List Bullet", "List Number"):
        set_font(styles[list_style], 10.5)
        styles[list_style].paragraph_format.space_after = Pt(4)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^]]+\]\([^)]+\))")


def add_inline(paragraph, value: str) -> None:
    for bit in INLINE.split(value):
        if not bit:
            continue
        if bit.startswith("**") and bit.endswith("**"):
            paragraph.add_run(bit[2:-2]).bold = True
        elif bit.startswith("*") and bit.endswith("*"):
            paragraph.add_run(bit[1:-1]).italic = True
        elif bit.startswith("`") and bit.endswith("`"):
            run = paragraph.add_run(bit[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif bit.startswith("[") and "](" in bit:
            label = bit[1:bit.index("](")]
            paragraph.add_run(label)
        else:
            paragraph.add_run(bit)


def set_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), "D9D9D9")


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def padding(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge in ("top", "left", "bottom", "right"):
        item = mar.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            mar.append(item)
        item.set(qn("w:w"), "90")
        item.set(qn("w:type"), "dxa")


def add_table(doc, rows: list[list[str]], landscape: bool) -> None:
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = False
    total = 9.5 if landscape else 7.0
    if cols == 5:
        proportions = [0.85, 1.1, 1.8, 1.1, 2.15]
    elif cols == 4:
        proportions = [1.05, 1.7, 1.35, 2.9]
    elif cols == 3:
        proportions = [2.3, 1.25, 3.45] if "Riesgo" in rows[0][0] else [0.9, 2.75, 3.35]
    elif cols == 2:
        proportions = [2.1, 4.9]
    else:
        proportions = [1] * cols
    factor = total / sum(proportions)
    widths = [Inches(x * factor) for x in proportions]
    for i, row in enumerate(rows):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        tr_pr.append(no_split)
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            padding(cell)
            set_border(cell)
            if i == 0:
                shade(cell, "DFEAF4")
            elif i % 2 == 0:
                shade(cell, "F7F9FC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, value)
            for run in p.runs:
                run.font.size = Pt(8.7 if cols >= 4 else 9.2)
                if i == 0:
                    run.bold = True
            if j == 0 and cols >= 3 and "Riesgo" not in rows[0][0]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def convert(md_path: Path, target: Path, landscape: bool) -> None:
    doc = Document()
    configure(doc, landscape)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_together = True
            for j, text in enumerate(block):
                if j:
                    p.add_run().add_break()
                run = p.add_run(text)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
            i += 1
            continue
        if not line:
            i += 1
            continue
        if md_path.name == "02-test-plan.md" and line == "## Estructura del proyecto":
            while i < len(lines) and lines[i].strip() != "## Resultado de la verificación":
                i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                candidate = lines[i].strip()
                cells = [c.strip() for c in candidate.strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", c or " ") for c in cells):
                    table_lines.append(cells)
                i += 1
            if table_lines:
                add_table(doc, table_lines, landscape)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_inline(p, line[2:])
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:])
            if md_path.name == "01-scrum-y-reparto.md" and line == "## Sprint sugerido de cinco jornadas":
                p.paragraph_format.page_break_before = True
            if md_path.name == "02-test-plan.md" and line == "## Matriz de trazabilidad resumida":
                p.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:])
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            add_inline(p, line)
        else:
            p = doc.add_paragraph(style="Normal")
            add_inline(p, line)
        i += 1
    doc.core_properties.title = lines[0].removeprefix("# ").strip()
    doc.core_properties.author = "Equipo QA"
    doc.save(target)


def main() -> None:
    parser = argparse.ArgumentParser(description="Regenera los documentos Word de entrega.")
    parser.add_argument("--source-prefix", help="Regenera solo fuentes cuyo nombre empiece con este prefijo.")
    args = parser.parse_args()
    OUT.mkdir(exist_ok=True)
    for source, name, landscape in JOBS:
        if args.source_prefix and not source.startswith(args.source_prefix):
            continue
        convert(ROOT / source, OUT / name, landscape)
        print(OUT / name)


if __name__ == "__main__":
    main()
